from collections.abc import Iterator
import csv
import contextvars
import difflib
import heapq
import os
import re
import stat
from contextlib import contextmanager
from datetime import datetime, timedelta
from pathlib import Path, PureWindowsPath

try:
    from rapidfuzz import fuzz
except ImportError:
    fuzz = None

_CURRENT_TASK_DIRECTORY: contextvars.ContextVar[Path | None] = contextvars.ContextVar(
    "hsx_task_output_directory", default=None
)


class Workspace:
    def __init__(self, root: Path, work_dir: str | Path = "Hsx Work"):
        self.root = Path(root).resolve()
        configured_work_dir = Path(work_dir)
        if configured_work_dir.is_absolute():
            unresolved_work_dir = configured_work_dir.absolute()
        else:
            unresolved_work_dir = (self.root / configured_work_dir).absolute()
        self._assert_inside(self.root, unresolved_work_dir)
        if unresolved_work_dir == self.root:
            raise ValueError("HSX_WORK_DIR must be a dedicated subdirectory of HSX_WORKSPACE.")
        current = unresolved_work_dir
        while current != self.root:
            if current.is_symlink():
                raise ValueError("HSX_WORK_DIR must not contain symlinks.")
            current = current.parent
        resolved_work_dir = unresolved_work_dir.resolve()
        self._assert_inside(self.root, resolved_work_dir)
        self.work_dir = resolved_work_dir

    @staticmethod
    def _assert_inside(root: Path, candidate: Path) -> None:
        if candidate != root and root not in candidate.parents:
            raise ValueError("Path must remain inside the workspace.")

    @staticmethod
    def _relative_path(raw_path: str | Path) -> Path:
        text = str(raw_path)
        windows_path = PureWindowsPath(text)
        path = Path(text)
        if path.is_absolute() or windows_path.is_absolute() or windows_path.drive:
            raise ValueError("Absolute paths are not allowed.")
        if ".." in path.parts or ".." in windows_path.parts:
            raise ValueError("Parent traversal is not allowed.")
        return path

    def resolve_read(self, raw_path: str | Path) -> Path:
        relative = self._relative_path(raw_path)
        candidate = (self.root / relative).resolve()
        self._assert_inside(self.root, candidate)
        return candidate

    def resolve_write(self, raw_path: str | Path) -> Path:
        relative = self._relative_path(raw_path)
        candidate = (self.work_dir / relative).resolve()
        self._assert_inside(self.work_dir, candidate)
        return candidate

    def unique_write_path(self, raw_path: str | Path) -> Path:
        candidate = self.resolve_write(raw_path)
        candidate.parent.mkdir(parents=True, exist_ok=True)
        if not candidate.exists():
            return candidate
        suffix = candidate.suffix
        stem = candidate.stem if suffix else candidate.name
        for index in range(1, 100_000):
            numbered = candidate.with_name(f"{stem} ({index}){suffix}")
            if not numbered.exists():
                return numbered
        raise FileExistsError("Could not find an unused output filename.")

    def create_task_dir(self, short_name: str, now: datetime | None = None) -> Path:
        if self.work_dir.resolve() != self.work_dir:
            raise ValueError("HSX_WORK_DIR changed through a symlink.")
        timestamp = (now or datetime.now()).strftime("%Y-%m-%d %H%M")
        safe_name = re.sub(
            r"(?i)\b(?:api[_ -]?key|token|password|secret)\s*[:=]\s*\S+",
            " ",
            short_name,
        )
        safe_name = re.sub(r"\b[A-Za-z0-9_-]{32,}\b", " ", safe_name)
        name = re.sub(r"[^A-Za-z -]+", " ", safe_name).strip()
        name = re.sub(r"\s+", " ", name)[:48].strip() or "task"
        return self._create_unique_directory(self.work_dir, f"{timestamp} {name}")

    @staticmethod
    def _create_unique_directory(parent: Path, name: str) -> Path:
        parent.mkdir(parents=True, exist_ok=True)
        for index in range(100_000):
            suffix = "" if index == 0 else f" ({index})"
            candidate = parent / f"{name}{suffix}"
            try:
                candidate.mkdir()
                return candidate.resolve()
            except FileExistsError:
                continue
        raise FileExistsError("Could not find an unused task directory.")

    def iter_files(self, directory: str | Path = ".") -> Iterator[Path]:
        target = self.resolve_read(directory)
        if not target.is_dir():
            raise ValueError("The requested workspace path is not a directory.")
        for current, directories, filenames in os.walk(
            target,
            followlinks=False,
            onerror=self._raise_walk_error,
        ):
            current_path = Path(current)
            safe_directories = []
            for name in directories:
                path = current_path / name
                attributes = getattr(os.lstat(path), "st_file_attributes", 0)
                if path.is_symlink() or (
                    os.name == "nt"
                    and attributes & getattr(stat, "FILE_ATTRIBUTE_REPARSE_POINT", 0)
                ):
                    continue
                try:
                    self._assert_inside(self.root, path.resolve())
                except ValueError:
                    continue
                safe_directories.append(name)
            directories[:] = safe_directories
            for name in filenames:
                path = current_path / name
                try:
                    resolved = path.resolve()
                    self._assert_inside(self.root, resolved)
                except (OSError, ValueError):
                    continue
                if resolved.is_file():
                    yield resolved

    @staticmethod
    def _raise_walk_error(error: OSError) -> None:
        raise error

    def files(self, directory: str | Path = ".") -> list[Path]:
        return list(self.iter_files(directory))

    def info(self, raw_path: str | Path) -> dict:
        path = self.resolve_read(raw_path)
        if not path.is_file():
            raise FileNotFoundError("Workspace file not found.")
        stat = path.stat()
        return {
            "name": path.name,
            "path": path.relative_to(self.root).as_posix(),
            "size": stat.st_size,
            "type": path.suffix.lower().lstrip(".") or "unknown",
            "modified": datetime.fromtimestamp(stat.st_mtime).astimezone().isoformat(),
        }

    def read_text(self, raw_path: str | Path, max_bytes: int = 5_000_000, max_chars: int = 100_000) -> str:
        path = self.resolve_read(raw_path)
        if not path.is_file():
            raise FileNotFoundError("Workspace file not found.")
        if path.stat().st_size > max_bytes:
            raise ValueError(f"File exceeds the {max_bytes}-byte read limit.")
        suffix = path.suffix.lower()
        if suffix in {".txt", ".md", ".json", ".log"}:
            text = path.read_text(encoding="utf-8")
        elif suffix == ".csv":
            with path.open(newline="", encoding="utf-8-sig") as file:
                text = "\n".join("\t".join(row) for row in csv.reader(file))
        elif suffix == ".docx":
            from docx import Document

            document = Document(path)
            text = "\n".join(
                [paragraph.text for paragraph in document.paragraphs]
                + ["\t".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
            )
        elif suffix == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                lines = []
                for sheet in workbook.worksheets:
                    lines.append(f"[{sheet.title}]")
                    for row in sheet.iter_rows(values_only=True):
                        lines.append("\t".join("" if value is None else str(value) for value in row))
                text = "\n".join(lines)
            finally:
                workbook.close()
        elif suffix == ".pdf":
            import pdfplumber

            with pdfplumber.open(path) as document:
                text = "\n".join(page.extract_text() or "" for page in document.pages)
        else:
            raise ValueError("Supported file types are txt, md, csv, docx, xlsx, and pdf.")
        return text[:max_chars]


def workspace_search(
    workspace: Workspace,
    query: str = "",
    directory: str = ".",
    file_type: str = "",
    recent_days: int | None = None,
    limit: int = 50,
) -> list[dict]:
    if limit < 1 or limit > 200:
        raise ValueError("Search limit must be between 1 and 200.")
    if recent_days is not None and recent_days < 0:
        raise ValueError("recent_days must not be negative.")
    requested_type = file_type.strip().lower().lstrip(".")
    query = query.strip()
    cutoff = (
        datetime.now().timestamp() - timedelta(days=recent_days).total_seconds()
        if recent_days is not None
        else None
    )
    best = []
    normalized_query = query.casefold()
    for path in workspace.iter_files(directory):
        if requested_type and path.suffix.lower().lstrip(".") != requested_type:
            continue
        stat = path.stat()
        if cutoff is not None and stat.st_mtime < cutoff:
            continue
        if query:
            score = (
                fuzz.WRatio(query, path.name)
                if fuzz is not None
                else max(
                    difflib.SequenceMatcher(None, normalized_query, path.name.casefold()).ratio() * 100,
                    difflib.SequenceMatcher(None, normalized_query, path.stem.casefold()).ratio() * 100,
                )
            )
            if score < 35:
                continue
        else:
            score = 0
        hit = (score, stat.st_mtime, str(path), path)
        if len(best) < limit:
            heapq.heappush(best, hit)
        elif hit[:3] > best[0][:3]:
            heapq.heapreplace(best, hit)
    best.sort(key=lambda hit: (hit[0], hit[1]), reverse=True)
    return [workspace.info(path.relative_to(workspace.root)) for _, _, _, path in best]


def current_task_dir(workspace: Workspace, task_name: str = "task") -> Path:
    return workspace.create_task_dir(task_name)


def task_output_path(workspace: Workspace, filename: str, task_name: str | None = None) -> Path:
    safe_name = PureWindowsPath(filename).name
    if safe_name in {"", ".", ".."}:
        raise ValueError("A valid output filename is required.")
    folder = _CURRENT_TASK_DIRECTORY.get()
    if folder is None:
        folder = workspace.create_task_dir(task_name or Path(safe_name).stem)
    relative = folder.resolve().relative_to(workspace.work_dir)
    return workspace.unique_write_path(relative / safe_name)


@contextmanager
def task_output_scope(workspace: Workspace, task_name: str, relative_directory: str = ""):
    if relative_directory:
        directory = workspace.resolve_write(relative_directory)
        if not directory.is_dir():
            raise FileNotFoundError("The task output directory does not exist.")
    else:
        directory = workspace.create_task_dir(task_name)
    token = _CURRENT_TASK_DIRECTORY.set(directory)
    try:
        yield directory
    finally:
        _CURRENT_TASK_DIRECTORY.reset(token)
