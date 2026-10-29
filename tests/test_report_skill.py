from pathlib import Path

from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas

from skills.report.skill import extract_template, make_pdf, read_data
from agent.workspace import Workspace


def test_report_extracts_docx_template_reads_data_and_creates_pdf(tmp_path: Path):
    root = tmp_path / "share"
    work = root / "Hsx Work"
    work.mkdir(parents=True)
    template = root / "template.docx"
    document = Document()
    document.add_heading("Quarterly report", level=1)
    document.add_paragraph("Prepared for the leadership team.")
    document.add_table(rows=2, cols=2)
    document.sections[0].header.paragraphs[0].text = "Company confidential"
    document.save(template)

    data_file = root / "source.xlsx"
    workbook = Workbook()
    workbook.active.append(["Metric", "Value"])
    workbook.active.append(["Revenue", 120])
    workbook.save(data_file)

    structure = extract_template(Workspace(root, work), "template.docx")
    data = read_data(Workspace(root, work), "source.xlsx")
    assert structure["sections"][0]["heading"] == "Quarterly report"
    assert structure["headers"] == ["Company confidential"]
    assert structure["tables"][0]["columns"] == 2
    assert data["tables"][0]["rows"][1] == ["Revenue", 120]

    output = make_pdf(
        Workspace(root, work),
        structure,
        {"title": "Quarterly report", "sections": [{"heading": "Results", "body": "Revenue was 120."}]},
        "result.pdf",
    )
    assert output["success"]
    assert "approximate" in output["fidelity"].lower()
    assert Path(output["output_path"]).is_file()
    assert Path(output["output_path"]).parent.parent == work


def test_pdf_template_extraction_and_read_data(tmp_path: Path):
    root = tmp_path / "share"
    work = root / "Hsx Work"
    work.mkdir(parents=True)
    source = root / "template.pdf"
    pdf = canvas.Canvas(str(source), pagesize=(612, 792))
    pdf.setFont("Helvetica-Bold", 20)
    pdf.drawString(72, 720, "Annual summary")
    pdf.setFont("Helvetica", 11)
    pdf.drawString(72, 690, "A reusable template.")
    pdf.save()

    workspace = Workspace(root, work)
    structure = extract_template(workspace, "template.pdf")
    data = read_data(workspace, "template.pdf")
    assert structure["source_type"] == "pdf"
    assert any(section["heading"] == "Annual summary" for section in structure["sections"])
    assert "A reusable template." in data["text"]
