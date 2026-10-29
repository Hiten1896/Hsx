from pathlib import Path

from docx import Document

from agent.workspace import Workspace, task_output_path
from config import WORKSPACE_DIR, WORK_DIR

WORKSPACE = Workspace(WORKSPACE_DIR, WORK_DIR)


def run(filename: str, title: str, sections: list) -> dict:
    path = task_output_path(WORKSPACE, f"{Path(filename).stem}.docx")
    document = Document()
    document.add_heading(title, 0)
    for section in sections:
        document.add_heading(section.get("heading", ""), level=1)
        document.add_paragraph(section.get("body", ""))
    with path.open("xb") as document_file:
        document.save(document_file)
    return {"success": True, "output_path": str(path)}